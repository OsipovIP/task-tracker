from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse, HttpResponse
from django.contrib import messages
from django.db.models import Q
from django.contrib.auth import get_user_model
from django.utils import timezone
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from .forms import TaskForm, TaskStatusForm, TaskPhotoForm, TaskCommentForm, TaskEditForm, TaskCompletionReportForm, CreateAndCloseTaskForm, RadioModelForm, RadioTypeForm, RadioAssignmentForm, RadioRepairForm
from .models import Task, OesObject, OesModel, TaskPhoto, TaskComment, TaskCompletionReport, ManualAsset, FaultCategory, RadioModel, RadioType, RadioAssignment, RadioRepair
import logging
from .models import IdleRecord, ModelTagConfig
from .clickhouse_client import ClickHouseClient
from collections import defaultdict

User = get_user_model()
logger = logging.getLogger(__name__) 


# 1. ВОССТАНОВЛЕННАЯ ФУНКЦИЯ: task_list
@login_required 
def task_list(request):
    """Отображает список всех задач, видимых пользователю (кроме выполненных)."""
    # Получаем все задачи, назначенные текущему пользователю, исключая выполненные
    # ITIL функционал отключен
    tasks = Task.objects.filter(
        assigned_to=request.user
    ).exclude(
        status=Task.STATUS_DONE
    ).order_by('-created_at') 

    context = {
        'tasks': tasks,
        'page_title': 'Мои задачи',
    }
    return render(request, 'tasks/task_list.html', context)


@login_required
def export_tasks_to_excel(request):
    """Экспорт задач в Excel файл"""
    # Получаем те же фильтры, что и в all_tasks
    status_filter = request.GET.get('status')
    asset_filter = request.GET.get('asset')
    assignee_filter = request.GET.get('assignee')
    date_from = request.GET.get('date_from')
    date_to = request.GET.get('date_to')
    category_filter = request.GET.get('category')
    
    # Если статус не указан и нет других фильтров, по умолчанию показываем "Выполнено"
    if status_filter is None and not asset_filter and not assignee_filter:
        status_filter = Task.STATUS_DONE
    
    # Базовый queryset с оптимизацией
    tasks = Task.objects.select_related('assigned_to', 'oes_object', 'oes_object__model', 'completion_report__fault_category').prefetch_related('additional_assignees').all()
    
    # Применяем фильтры
    if status_filter:
        tasks = tasks.filter(status=status_filter)
    if asset_filter:
        tasks = tasks.filter(oes_object__name__icontains=asset_filter)
    if assignee_filter:
        tasks = tasks.filter(
            Q(assigned_to__username__icontains=assignee_filter) |
            Q(assigned_to__first_name__icontains=assignee_filter) |
            Q(assigned_to__last_name__icontains=assignee_filter) |
            Q(additional_assignees__username__icontains=assignee_filter) |
            Q(additional_assignees__first_name__icontains=assignee_filter) |
            Q(additional_assignees__last_name__icontains=assignee_filter)
        ).distinct()
    if date_from:
        tasks = tasks.filter(closed_at__date__gte=date_from)
    if date_to:
        tasks = tasks.filter(closed_at__date__lte=date_to)
    if category_filter:
        tasks = tasks.filter(completion_report__fault_category_id=category_filter)
    
    tasks = tasks.order_by('-priority', 'created_at')
    
    # Создаем Excel файл
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Задачи"
    
    # Заголовки
    headers = [
        'ID', 'Актив', 'Статус', 'Исполнитель', 'Доп. исполнитель', 
        'Создана', 'Закрыта', 'Категория', 'Выполненная работа'
    ]
    
    # Стили для заголовков
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    # Записываем заголовки
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
    
    # Записываем данные
    for row, task in enumerate(tasks, 2):
        # ID
        ws.cell(row=row, column=1, value=f"#{task.id}")
        
        # Актив
        if task.oes_object:
            asset_name = task.oes_object.name
            if task.oes_object.model:
                asset_name += f" ({task.oes_object.model.name})"
            ws.cell(row=row, column=2, value=asset_name)
        else:
            ws.cell(row=row, column=2, value="Не указан")
        
        # Статус
        ws.cell(row=row, column=3, value=task.get_status_display())
        
        # Исполнитель
        if task.assigned_to:
            if task.assigned_to.first_name or task.assigned_to.last_name:
                assignee = f"{task.assigned_to.first_name} {task.assigned_to.last_name}".strip()
            else:
                assignee = task.assigned_to.username
        elif task.assigned_group:
            assignee = f"Пул ({task.assigned_group.name})"
        else:
            assignee = "Не назначена"
        ws.cell(row=row, column=4, value=assignee)
        
        # Дополнительные исполнители
        if task.additional_assignees.all():
            additional_assignees = []
            for user in task.additional_assignees.all():
                if user.first_name or user.last_name:
                    additional_assignees.append(f"{user.first_name} {user.last_name}".strip())
                else:
                    additional_assignees.append(user.username)
            ws.cell(row=row, column=5, value=", ".join(additional_assignees))
        else:
            ws.cell(row=row, column=5, value="-")
        
        # Создана
        ws.cell(row=row, column=6, value=task.created_at.strftime("%d.%m.%Y"))
        
        # Закрыта
        if task.closed_at:
            ws.cell(row=row, column=7, value=task.closed_at.strftime("%d.%m.%Y"))
        else:
            ws.cell(row=row, column=7, value="-")
        
        # Категория
        if hasattr(task, 'completion_report') and task.completion_report and task.completion_report.fault_category:
            ws.cell(row=row, column=8, value=task.completion_report.fault_category.name)
        else:
            ws.cell(row=row, column=8, value="-")
        
        # Выполненная работа
        if hasattr(task, 'completion_report') and task.completion_report and task.completion_report.work_performed:
            ws.cell(row=row, column=9, value=task.completion_report.work_performed)
        else:
            ws.cell(row=row, column=9, value="-")
    
    # Автоподбор ширины колонок
    for col in range(1, len(headers) + 1):
        column_letter = get_column_letter(col)
        max_length = 0
        for row in range(1, ws.max_row + 1):
            cell_value = ws[f"{column_letter}{row}"].value
            if cell_value:
                max_length = max(max_length, len(str(cell_value)))
        ws.column_dimensions[column_letter].width = min(max_length + 2, 50)
    
    # Создаем HTTP ответ
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="tasks_export.xlsx"'
    
    wb.save(response)
    return response
# --------------------------------------------------------------------------


# 2. ВАША ФУНКЦИЯ: task_create
@login_required # Только для авторизованных пользователей
def task_create(request):
    if request.method == 'POST':
        form = TaskForm(request.POST)
        
        if form.is_valid():
            # Получаем ID объекта из скрытого поля
            oes_object_id = request.POST.get('oes_object_id')
            
            if oes_object_id:
                try:
                    oes_object = OesObject.objects.get(id=oes_object_id)
                    
                    # Создаем задачу с данными из формы
                    new_task = form.save(commit=False)
                    new_task.oes_object = oes_object
                    new_task.reporter = request.user 
                    
                    # Если исполнитель не указан, назначаем на создателя
                    if not new_task.assigned_to:
                        new_task.assigned_to = request.user
                    
                    new_task.save(update_fields=['assigned_to', 'oes_object', 'updated_at'])
                    
                    messages.success(request, 'Задача успешно создана!')
                    return redirect('task_detail', pk=new_task.pk) 
                    
                except OesObject.DoesNotExist:
                    messages.error(request, 'Выбранный объект не найден')
            else:
                messages.error(request, 'Пожалуйста, выберите объект')
        else:
            messages.error(request, 'Пожалуйста, исправьте ошибки в форме')
    else:
        form = TaskForm()
        
    context = {
        'form': form, 
        'page_title': 'Создание новой задачи',
    }
    return render(request, 'task_form.html', context)


# 3. ФУНКЦИЯ: task_detail
@login_required
def task_detail(request, pk):
    """Отображает детали конкретной задачи."""
    task = get_object_or_404(Task, pk=pk)
    
    # Определяем, может ли пользователь редактировать задачу
    # Редактирование доступно ТОЛЬКО менеджерам
    user_is_assigned = request.user.groups.filter(name='Менеджеры').exists()
    
    # Определяем, может ли пользователь работать с задачей (изменять статус, загружать фото)
    # ITIL функционал отключен - убрана логика для ITIL задач без исполнителя
    user_can_work = (
        task.assigned_to == request.user or 
        task.reporter == request.user
        # or (task.external_source == 'ITIL' and not task.assigned_to)  # ITIL отключен
    )
    
    # Получаем фотографии и комментарии (исключаем пустые фото)
    photos = TaskPhoto.objects.filter(task=task).exclude(image='').order_by('-uploaded_at')
    comments = TaskComment.objects.filter(task=task).order_by('created_at')
    
    # Обработка POST запросов
    if request.method == 'POST':
        # Изменение статуса
        if 'status_submit' in request.POST:
            if user_can_work:
                new_status = request.POST.get('status')
                if new_status in [choice[0] for choice in Task.STATUS_CHOICES]:
                    task.status = new_status
                    if new_status == Task.STATUS_DONE:
                        task.closed_at = timezone.now()
                        task.save(update_fields=['status', 'closed_at', 'updated_at'])
                    else:
                        task.save(update_fields=['status', 'updated_at'])
                    messages.success(request, 'Статус задачи обновлен!')
                else:
                    messages.error(request, 'Неверный статус!')
            else:
                messages.error(request, 'Только назначенный исполнитель может изменять статус!')
        
        # Загрузка фото
        elif 'photo_submit' in request.POST:
            if user_can_work:
                print(f"DEBUG Photo: FILES = {request.FILES}")
                print(f"DEBUG Photo: POST = {request.POST}")
                
                photo_form = TaskPhotoForm(request.POST, request.FILES)
                print(f"DEBUG Photo: Form valid = {photo_form.is_valid()}")
                
                if photo_form.is_valid():
                    photo = photo_form.save(commit=False)
                    photo.task = task
                    photo.uploaded_by = request.user
                    print(f"DEBUG Photo: Before save, image = {photo.image}")
                    photo.save()
                    print(f"DEBUG Photo: After save, image.name = {photo.image.name if photo.image else 'EMPTY'}")
                    messages.success(request, 'Фото успешно загружено!')
                else:
                    print(f"DEBUG Photo: Form errors = {photo_form.errors}")
                    messages.error(request, f'Ошибка при загрузке фото: {photo_form.errors}')
            else:
                messages.error(request, 'Только назначенный исполнитель может загружать фото!')
        
        # Добавление комментария
        elif 'comment_submit' in request.POST:
            comment_text = request.POST.get('text', '').strip()
            if comment_text:
                TaskComment.objects.create(
                    task=task,
                    author=request.user,
                    text=comment_text
                )
                messages.success(request, 'Комментарий добавлен!')
            else:
                messages.error(request, 'Комментарий не может быть пустым!')
        
        # Редактирование задачи (для менеджеров)
        elif 'edit_submit' in request.POST:
            if user_is_assigned:
                print(f"DEBUG: Получены данные формы: {request.POST}")
                print(f"DEBUG: Текущая задача: {task.title}")
                
                # Получаем ID объекта из скрытого поля перед валидацией формы
                oes_object_id = request.POST.get('oes_object_id')
                print(f"DEBUG: oes_object_id из скрытого поля: {oes_object_id}")
                
                if oes_object_id:
                    try:
                        oes_object = OesObject.objects.get(id=oes_object_id)
                        # Временно устанавливаем объект для валидации
                        task.oes_object = oes_object
                        print(f"DEBUG: Установлен объект: {oes_object.name}")
                    except OesObject.DoesNotExist:
                        messages.error(request, 'Выбранный объект не найден')
                        oes_object_id = None
                
                edit_form = TaskEditForm(request.POST, instance=task)
                print(f"DEBUG: Форма валидна: {edit_form.is_valid()}")
                if not edit_form.is_valid():
                    print(f"DEBUG: Ошибки формы: {edit_form.errors}")
                
                if edit_form.is_valid():
                    # Сохраняем форму
                    saved_task = edit_form.save()
                    print(f"DEBUG: Задача сохранена: {saved_task.title}")
                    print(f"DEBUG: Объект после сохранения: {saved_task.oes_object}")
                    messages.success(request, f'Задача #{task.id} успешно обновлена!')
                else:
                    messages.error(request, f'Ошибка при обновлении задачи: {edit_form.errors}')
            else:
                messages.error(request, 'У вас нет прав для редактирования этой задачи!')
        
        # Закрытие задачи с отчетом
        elif 'complete_submit' in request.POST:
            if user_can_work:
                completion_form = TaskCompletionReportForm(request.POST, task=task)
                if completion_form.is_valid():
                    # Проверяем, существует ли уже отчет для этой задачи
                    try:
                        report = task.completion_report
                        # Обновляем существующий отчет
                        report.completed_by = request.user
                        report.fault_category = completion_form.cleaned_data.get('fault_category')
                        report.work_performed = completion_form.cleaned_data.get('work_performed')
                        report.root_cause = completion_form.cleaned_data.get('root_cause')
                        report.preventive_measures = completion_form.cleaned_data.get('preventive_measures')
                        report.time_spent = completion_form.cleaned_data.get('time_spent')
                        report.save()
                    except TaskCompletionReport.DoesNotExist:
                        # Создаем новый отчет
                        report = completion_form.save(commit=False)
                        report.task = task
                        report.completed_by = request.user
                        report.save()
                    
                    # Сохраняем дополнительных исполнителей
                    additional_assignees = completion_form.cleaned_data.get('additional_assignees')
                    if additional_assignees:
                        task.additional_assignees.set(additional_assignees)
                    else:
                        task.additional_assignees.clear()
                    
                    # Переводим задачу в статус "Проверка" для проверки отчета
                    task.status = Task.STATUS_CHECKING
                    task.save(update_fields=['status', 'updated_at'])
                    
                    messages.success(request, 'Задача переведена на проверку!')
                else:
                    messages.error(request, 'Ошибка при создании отчета!')
            else:
                messages.error(request, 'У вас нет прав для закрытия этой задачи!')
        
        return redirect('task_detail', pk=task.pk)
    
    # Создаем формы
    status_form = TaskStatusForm(initial={'status': task.status})
    photo_form = TaskPhotoForm()
    comment_form = TaskCommentForm()
    edit_form = TaskEditForm(instance=task)
    completion_form = TaskCompletionReportForm(task=task)
    
    # Проверяем, есть ли уже отчет о закрытии
    completion_report = getattr(task, 'completion_report', None)
    
    context = {
        'task': task,
        'user_is_assigned': user_is_assigned,  # Только для редактирования (менеджеры)
        'user_can_work': user_can_work,       # Для работы с задачей (исполнители)
        'photos': photos,
        'comments': comments,
        'status_form': status_form,
        'photo_form': photo_form,
        'comment_form': comment_form,
        'edit_form': edit_form,
        'completion_form': completion_form,
        'completion_report': completion_report,
        'page_title': f'Задача #{task.id}',
    }
    return render(request, 'tasks/task_detail.html', context)


# 4. ФУНКЦИЯ: task_claim - захват задачи из пула
@login_required
def task_claim(request, pk):
    """Позволяет пользователю взять задачу на себя."""
    task = get_object_or_404(Task, pk=pk)
    
    # Назначаем текущего пользователя исполнителем
    task.assigned_to = request.user
    task.status = Task.STATUS_TODO
    task.save(update_fields=['assigned_to', 'status', 'updated_at'])
    
    return redirect('task_detail', pk=task.pk)


# 5. ФУНКЦИЯ: task_confirm_receipt - подтверждение получения задачи
@login_required
def task_confirm_receipt(request, pk):
    """Подтверждает, что пользователь получил задачу."""
    from django.utils import timezone
    
    task = get_object_or_404(Task, pk=pk)
    
    # Устанавливаем время подтверждения
    task.acknowledged_at = timezone.now()
    task.save(update_fields=['acknowledged_at', 'updated_at'])
    
    return redirect('task_detail', pk=task.pk)


# 6. ФУНКЦИЯ: load_objects - AJAX загрузка объектов по модели
from django.http import JsonResponse

def load_objects(request):
    """AJAX endpoint для поиска OesObject и ManualAsset по тексту."""
    search_query = request.GET.get('search')
    model_id = request.GET.get('model_id')  # Для обратной совместимости
    
    if search_query and len(search_query) >= 2:
        # Улучшенный поиск: по названию объекта, модели и категории
        from django.db.models import Q
        
        formatted_objects = []
        
        # Поиск OES объектов
        oes_objects = OesObject.objects.filter(
            Q(name__icontains=search_query) |
            Q(model__name__icontains=search_query) |
            Q(model__category__name__icontains=search_query)
        ).select_related('model', 'model__category').distinct()[:10]
        
        for obj in oes_objects:
            # Исправление: проверка на наличие category перед доступом
            if obj.model and obj.model.category:
                model_name = f"{obj.model.category.name} / {obj.model.name}"
            elif obj.model:
                model_name = obj.model.name
            else:
                model_name = "Неизвестная модель"
            
            formatted_objects.append({
                'id': obj.id,
                'name': obj.name,
                'model_name': model_name,
                'type': 'oes',
                'display_name': f"{obj.name} ({model_name})"
            })
            
        
        # Поиск ручных активов
        manual_assets = ManualAsset.objects.filter(
            Q(name__icontains=search_query) |
            Q(asset_type__icontains=search_query) |
            Q(model__icontains=search_query) |
            Q(description__icontains=search_query),
            is_active=True
        ).distinct()[:10]
        
        for asset in manual_assets:
            asset_info = asset.asset_type
            if asset.model:
                asset_info += f" / {asset.model}"
            formatted_objects.append({
                'id': asset.id,
                'name': asset.name,
                'model_name': asset_info,
                'type': 'manual',
                'display_name': f"{asset.name} ({asset_info}) [Ручной]"
            })
        
        # Сортируем по релевантности (сначала точные совпадения в названии)
        formatted_objects.sort(key=lambda x: (
            0 if search_query.lower() in x['name'].lower() else 1,
            x['name']
        ))
        
        return JsonResponse({'objects': formatted_objects[:15]})
    
    # Старая логика для model_id (обратная совместимость)
    elif model_id:
        objects = OesObject.objects.filter(model_id=model_id).values('id', 'name')
        return JsonResponse({'objects': list(objects)})
    
    return JsonResponse({'objects': []})


# 7. ФУНКЦИЯ: task_pool - задачи со статусом "к выполнению"
@login_required
def task_pool(request):
    """Отображает наряд - все задачи со статусом 'к выполнению'."""
    
    # Отладка
    logger.info(f"task_pool: method={request.method}, POST data={request.POST}")
    
    # Обработка POST запроса (отправка задачи в Telegram)
    if request.method == 'POST' and 'send_task_telegram' in request.POST:
        task_id = request.POST.get('send_task_telegram')
        try:
            import subprocess
            logger.info(f"Running send_single_task_to_telegram command for task {task_id}")
            result = subprocess.run(
                ['python', 'manage.py', 'send_single_task_to_telegram', str(task_id)],
                cwd='/usr/src/app',
                capture_output=True,
                text=True,
                timeout=10
            )
            if result.returncode == 0:
                messages.success(request, f'Задача #{task_id} отправлена в Telegram!')
            else:
                messages.error(request, f'Ошибка при отправке задачи #{task_id}: {result.stderr}')
        except Exception as e:
            messages.error(request, f'Ошибка при отправке задачи #{task_id}: {str(e)}')
            logger.error(f"Failed to send single task from task_pool: {e}", exc_info=True)
        
        return redirect('task_pool')
    
    tasks = Task.objects.filter(status=Task.STATUS_TODO).order_by('-priority', 'created_at')
    
    context = {
        'tasks': tasks,
        'page_title': 'Наряд',
    }
    return render(request, 'tasks/task_pool.html', context)


# 8. ФУНКЦИЯ: all_tasks - все задачи с фильтром по статусу
@login_required
def all_tasks(request):
    """Отображает все задачи с возможностью фильтрации по статусу, активу и назначенному пользователю."""
    from django.contrib.auth import get_user_model
    
    # Получаем параметры фильтрации из GET запроса
    status_filter = request.GET.get('status')
    asset_filter = request.GET.get('asset')
    assignee_filter = request.GET.get('assignee')
    date_from = request.GET.get('date_from')
    date_to = request.GET.get('date_to')
    category_filter = request.GET.get('category')
    
    # Если статус не указан и нет других фильтров, по умолчанию показываем "Выполнено"
    if status_filter is None and not asset_filter and not assignee_filter:
        status_filter = Task.STATUS_DONE
    
    # Базовый queryset с оптимизацией
    tasks = Task.objects.select_related('assigned_to', 'oes_object', 'oes_object__model', 'completion_report__fault_category').prefetch_related('additional_assignees').all()
    
    # Применяем фильтр по статусу, если он указан
    if status_filter:
        tasks = tasks.filter(status=status_filter)
    
    # Применяем фильтр по активу, если он указан
    if asset_filter:
        tasks = tasks.filter(oes_object__name__icontains=asset_filter)
    
    # Применяем фильтр по назначенному пользователю, если он указан
    if assignee_filter:
        tasks = tasks.filter(
            Q(assigned_to__username__icontains=assignee_filter) |
            Q(assigned_to__first_name__icontains=assignee_filter) |
            Q(assigned_to__last_name__icontains=assignee_filter) |
            Q(additional_assignees__username__icontains=assignee_filter) |
            Q(additional_assignees__first_name__icontains=assignee_filter) |
            Q(additional_assignees__last_name__icontains=assignee_filter)
        ).distinct()
    
    # Применяем фильтр по дате закрытия, если указан
    if date_from:
        tasks = tasks.filter(closed_at__date__gte=date_from)
    if date_to:
        tasks = tasks.filter(closed_at__date__lte=date_to)
    
    # Применяем фильтр по категории, если указан
    if category_filter:
        tasks = tasks.filter(completion_report__fault_category_id=category_filter)
    
    tasks = tasks.order_by('-priority', 'created_at')
    
    # Получаем все возможные статусы для фильтра
    status_choices = Task.STATUS_CHOICES
    
    # Получаем уникальные активы для автокомплита
    assets = OesObject.objects.values_list('name', flat=True).distinct().order_by('name')
    
    # Получаем пользователей для автокомплита
    User = get_user_model()
    users = User.objects.values_list('username', flat=True).distinct().order_by('username')
    
    # Получаем категории неисправностей
    from .models import FaultCategory
    categories = FaultCategory.objects.filter(is_active=True).order_by('name')
    
    context = {
        'tasks': tasks,
        'status_choices': status_choices,
        'current_status': status_filter,
        'current_asset': asset_filter,
        'current_assignee': assignee_filter,
        'current_date_from': date_from,
        'current_date_to': date_to,
        'current_category': category_filter,
        'assets': assets,
        'users': users,
        'categories': categories,
        'page_title': 'Все задачи',
    }
    return render(request, 'tasks/all_tasks.html', context)


# 9. ФУНКЦИЯ: work_order_management - формирование наряда (только для менеджеров)
@login_required
def work_order_management(request):
    """Быстрое редактирование задач для формирования наряда. Доступно только менеджерам."""
    from django.contrib.auth import get_user_model
    from django.contrib.auth.models import Group
    
    # Проверяем, что пользователь в группе "Менеджеры"
    if not request.user.groups.filter(name='Менеджеры').exists():
        messages.error(request, 'Доступ запрещен. Эта страница доступна только менеджерам.')
        return redirect('task_list')
    
    User = get_user_model()
    
    # Обработка POST запроса (сохранение изменений или создание новой задачи)
    if request.method == 'POST':
        # Отправка наряда в Telegram (проверяем первым)
        if 'send_to_telegram' in request.POST:
            logger.info("Send to Telegram button pressed")
            try:
                from django.core.management import call_command
                from io import StringIO
                
                logger.info("Running send_work_order_to_telegram command")
                out = StringIO()
                call_command('send_work_order_to_telegram', stdout=out)
                logger.info(f"Command output: {out.getvalue()}")
                messages.success(request, '✅ Наряд отправлен в Telegram!')
            except Exception as e:
                logger.error(f"Failed to send work order to telegram: {e}", exc_info=True)
                messages.error(request, f'Ошибка при отправке наряда в Telegram: {str(e)}')
            
            return redirect('work_order_management')
        
        create_new_task = request.POST.get('create_new_task')
        create_and_send_telegram = request.POST.get('create_and_send_telegram')
        
        # Создание новой задачи
        if create_new_task:
            try:
                # Получаем данные из формы
                title = request.POST.get('title', '').strip()
                description = request.POST.get('description', '').strip()
                status = request.POST.get('status', Task.STATUS_WAITING)
                assigned_to_id = request.POST.get('assigned_to')
                due_date = request.POST.get('due_date') or None
                itil_number = request.POST.get('itil_number', '').strip() or None
                oes_object_id = request.POST.get('oes_object_id')
                
                if not title:
                    messages.error(request, 'Заголовок задачи обязателен!')
                    return redirect('work_order_management')
                
                # Создаем новую задачу
                new_task = Task.objects.create(
                    title=title,
                    description=description,
                    status=status,
                    reporter=request.user,
                    due_date=due_date,
                    itil_number=itil_number
                )
                
                # Назначаем исполнителя
                if assigned_to_id:
                    try:
                        assigned_user = User.objects.get(id=assigned_to_id)
                        new_task.assigned_to = assigned_user
                    except User.DoesNotExist:
                        pass
                
                # Привязываем актив (OES или ручной)
                manual_asset_id = request.POST.get('manual_asset_id')
                if oes_object_id:
                    try:
                        oes_object = OesObject.objects.get(id=oes_object_id)
                        new_task.oes_object = oes_object
                        new_task.manual_asset = None  # Очищаем ручной актив, если выбран OES
                    except OesObject.DoesNotExist:
                        pass
                elif manual_asset_id:
                    try:
                        manual_asset = ManualAsset.objects.get(id=manual_asset_id)
                        new_task.manual_asset = manual_asset
                        new_task.oes_object = None  # Очищаем OES объект, если выбран ручной актив
                    except ManualAsset.DoesNotExist:
                        pass
                
                update_fields = ['updated_at']
                if assigned_to_id:
                    update_fields.append('assigned_to')
                if oes_object_id:
                    update_fields.append('oes_object')
                if manual_asset_id:
                    update_fields.append('manual_asset')
                new_task.save(update_fields=update_fields)
                
                # Отправляем уведомление в Telegram, если задача назначена
                if new_task.assigned_to:
                    logger.info(f"Trying to send telegram notification for task #{new_task.id} to user {new_task.assigned_to.id}")
                    try:
                        import subprocess
                        subprocess.Popen([
                            'python', 'manage.py', 'send_telegram_notification',
                            str(new_task.assigned_to.id), str(new_task.id)
                        ], cwd='/usr/src/app')
                        logger.info(f"Telegram notification command triggered")
                    except Exception as e:
                        logger.error(f"Failed to send telegram notification: {e}", exc_info=True)
                
                # Если нужно отправить в Telegram после создания
                if create_and_send_telegram:
                    try:
                        from django.core.management import call_command
                        from io import StringIO
                        logger.info(f"Running send_single_task_to_telegram command for new task {new_task.id}")
                        out = StringIO()
                        call_command('send_single_task_to_telegram', str(new_task.id), stdout=out)
                        logger.info(f"Command output: {out.getvalue()}")
                        messages.success(request, f'Задача #{new_task.id} "{title}" создана и отправлена в Telegram!')
                    except Exception as e:
                        messages.success(request, f'Задача #{new_task.id} "{title}" создана, но ошибка при отправке: {str(e)}')
                        logger.error(f"Failed to send single task after creation: {e}", exc_info=True)
                else:
                    messages.success(request, f'Задача #{new_task.id} "{title}" успешно создана!')
                
            except Exception as e:
                messages.error(request, f'Ошибка при создании задачи: {str(e)}')
            
            return redirect('work_order_management')
        
        # Проверяем, нужно ли отправить в Telegram после сохранения
        save_and_send_telegram = request.POST.get('save_and_send_telegram')
        
        # Редактирование существующей задачи
        task_id = request.POST.get('task_id')
        if task_id:
            try:
                task = Task.objects.get(id=task_id)
                
                # Обновляем поля задачи
                old_status = task.status
                task.title = request.POST.get('title', task.title)
                task.description = request.POST.get('description', task.description)
                task.status = request.POST.get('status', task.status)
                task.priority = request.POST.get('priority', task.priority)
                task.due_date = request.POST.get('due_date') or None
                task.itil_number = request.POST.get('itil_number') or None
                
                # Обрабатываем closed_at при изменении статуса
                if task.status == Task.STATUS_DONE:
                    if not task.closed_at:
                        task.closed_at = timezone.now()
                else:
                    # Очищаем дату закрытия для других статусов
                    task.closed_at = None
                
                # Обновляем актив (OES или ручной)
                oes_object_id = request.POST.get('oes_object_id')
                manual_asset_id = request.POST.get('manual_asset_id')
                if oes_object_id:
                    try:
                        task.oes_object = OesObject.objects.get(id=oes_object_id)
                        task.manual_asset = None  # Очищаем ручной актив, если выбран OES
                    except OesObject.DoesNotExist:
                        pass
                elif manual_asset_id:
                    try:
                        task.manual_asset = ManualAsset.objects.get(id=manual_asset_id)
                        task.oes_object = None  # Очищаем OES объект, если выбран ручной актив
                    except ManualAsset.DoesNotExist:
                        pass
                elif not oes_object_id and not manual_asset_id:
                    # Если оба поля пустые, очищаем оба актива
                    task.oes_object = None
                    task.manual_asset = None
                
                # Обновляем назначенного пользователя
                old_assignee = task.assigned_to
                assigned_to_id = request.POST.get('assigned_to')
                if assigned_to_id:
                    try:
                        new_assignee = User.objects.get(id=assigned_to_id)
                        task.assigned_to = new_assignee
                        
                        # Отправляем уведомление, если исполнитель изменился
                        if old_assignee != new_assignee:
                            try:
                                import subprocess
                                subprocess.Popen([
                                    'python', 'manage.py', 'send_telegram_notification',
                                    str(new_assignee.id), str(task.id)
                                ], cwd='/usr/src/app')
                                logger.info(f"Telegram notification command triggered for user {new_assignee.id}")
                            except Exception as e:
                                logger.warning(f"Failed to send telegram notification: {e}")
                    except User.DoesNotExist:
                        pass
                elif request.POST.get('assigned_to') == '':
                    task.assigned_to = None
                
                # Формируем список полей для обновления
                update_fields = ['title', 'description', 'status', 'priority', 'due_date', 'itil_number', 'assigned_to', 'updated_at']
                
                # Добавляем closed_at, если статус изменился
                if old_status != task.status:
                    update_fields.append('closed_at')
                
                if oes_object_id or task.oes_object is None:
                    update_fields.append('oes_object')
                if manual_asset_id or task.manual_asset is None:
                    update_fields.append('manual_asset')
                
                task.save(update_fields=update_fields)
                
                # Если нужно отправить в Telegram после сохранения
                if save_and_send_telegram:
                    try:
                        from django.core.management import call_command
                        from io import StringIO
                        logger.info(f"Running send_single_task_to_telegram command for task {task.id}")
                        out = StringIO()
                        call_command('send_single_task_to_telegram', str(task.id), stdout=out)
                        logger.info(f"Command output: {out.getvalue()}")
                        messages.success(request, f'Задача #{task.id} обновлена и отправлена в Telegram!')
                    except Exception as e:
                        messages.success(request, f'Задача #{task.id} обновлена, но ошибка при отправке: {str(e)}')
                        logger.error(f"Failed to send single task after update: {e}", exc_info=True)
                else:
                    messages.success(request, f'Задача #{task.id} успешно обновлена!')
                
            except Task.DoesNotExist:
                messages.error(request, 'Задача не найдена')
        
        return redirect('work_order_management')
    
    # Получаем все задачи со статусом "Ожидание"
    tasks = Task.objects.filter(status=Task.STATUS_WAITING).order_by('-created_at')
    
    # Получаем всех пользователей для выбора исполнителя
    users = User.objects.filter(is_active=True).order_by('username')
    
    # Получаем все статусы
    status_choices = Task.STATUS_CHOICES
    
    # Получаем активы для автокомплита
    assets = OesObject.objects.all().select_related('model').order_by('name')
    
    context = {
        'tasks': tasks,
        'users': users,
        'status_choices': status_choices,
        'assets': assets,
        'page_title': 'Формирование наряда',
    }
    return render(request, 'tasks/work_order_management.html', context)


@login_required
def tasks_for_checking(request):
    """Отображает список всех задач со статусом 'Проверка'."""
    
    # Обработка POST запросов (изменение статуса задачи)
    if request.method == 'POST':
        task_id = request.POST.get('task_id')
        new_status = request.POST.get('new_status')
        
        if task_id and new_status:
            try:
                task = Task.objects.get(id=task_id)
                
                # Проверяем права (только менеджеры могут менять статус)
                if request.user.groups.filter(name='Менеджеры').exists():
                    old_status = task.get_status_display()
                    task.status = new_status
                    
                    if new_status == Task.STATUS_DONE:
                        task.closed_at = timezone.now()
                        task.save(update_fields=['status', 'closed_at', 'updated_at'])
                        messages.success(request, f'Задача #{task.id} подтверждена и переведена в статус "Выполнено"!')
                    elif new_status == Task.STATUS_TODO:
                        # Очищаем дату закрытия при возврате в работу
                        task.closed_at = None
                        task.save(update_fields=['status', 'closed_at', 'updated_at'])
                        messages.warning(request, f'Задача #{task.id} возвращена в работу.')
                    else:
                        # Для других статусов тоже очищаем дату закрытия
                        if task.closed_at:
                            task.closed_at = None
                            task.save(update_fields=['status', 'closed_at', 'updated_at'])
                        else:
                            task.save(update_fields=['status', 'updated_at'])
                        messages.info(request, f'Статус задачи #{task.id} изменен.')
                else:
                    messages.error(request, 'У вас нет прав для изменения статуса задачи!')
            except Task.DoesNotExist:
                messages.error(request, 'Задача не найдена!')
        
        return redirect('tasks_for_checking')
    
    # Получаем все задачи со статусом "Проверка"
    tasks = Task.objects.filter(
        status=Task.STATUS_CHECKING
    ).order_by('-created_at')
    
    print(f"DEBUG tasks_for_checking: Найдено задач на проверке: {tasks.count()}")
    for task in tasks:
        print(f"DEBUG: ID={task.id}, Заголовок={task.title}, Статус={task.status}")
    
    # Получаем всех пользователей для фильтрации
    users = User.objects.all().order_by('username')
    
    # Статусы для фильтрации
    status_choices = Task.STATUS_CHOICES
    
    context = {
        'tasks': tasks,
        'users': users,
        'status_choices': status_choices,
        'page_title': 'Задачи на проверке',
    }
    return render(request, 'tasks/tasks_for_checking.html', context)


@login_required
def create_and_close_task(request):
    """Создание задачи с автоматическим закрытием для инженеров."""
    if request.method == 'POST':
        form = CreateAndCloseTaskForm(request.POST)
        
        if form.is_valid():
            # Создаем задачу
            task = form.save(commit=False)
            
            # Устанавливаем исполнителя на текущего пользователя
            task.assigned_to = request.user
            task.reporter = request.user
            
            # Устанавливаем актив (OES или ручной) из скрытых полей
            oes_object_id = request.POST.get('oes_object_id')
            manual_asset_id = request.POST.get('manual_asset_id')
            
            if oes_object_id:
                try:
                    task.oes_object = OesObject.objects.get(id=oes_object_id)
                    task.manual_asset = None  # Очищаем ручной актив, если выбран OES
                except OesObject.DoesNotExist:
                    pass
            elif manual_asset_id:
                try:
                    task.manual_asset = ManualAsset.objects.get(id=manual_asset_id)
                    task.oes_object = None  # Очищаем OES объект, если выбран ручной актив
                except ManualAsset.DoesNotExist:
                    pass
            
            # Устанавливаем статус "Выполнено"
            task.status = Task.STATUS_DONE
            task.closed_at = timezone.now()
            
            # Сохраняем задачу (первое сохранение - без update_fields, так как объект новый)
            task.save()
            
            # Создаем отчет о закрытии
            TaskCompletionReport.objects.create(
                task=task,
                completed_by=request.user,
                fault_category=form.cleaned_data.get('fault_category'),
                work_performed=form.cleaned_data.get('work_performed', ''),
                time_spent=None  # Можно добавить поле для времени
            )
            
            messages.success(request, f'Задача #{task.id} успешно создана и закрыта!')
            return redirect('all_tasks')
    else:
        form = CreateAndCloseTaskForm()
    
    context = {
        'form': form,
        'page_title': 'Создать и закрыть задачу',
    }
    return render(request, 'tasks/create_and_close_task.html', context)


# Views для учета раций
@login_required
def radio_models_list(request):
    """Список моделей раций."""
    models_list = RadioModel.objects.all().order_by('name')
    active_count = models_list.filter(is_active=True).count()
    
    context = {
        'models': models_list,
        'active_count': active_count,
        'page_title': 'Справочник моделей раций',
    }
    return render(request, 'tasks/radio_models_list.html', context)


@login_required
def radio_model_add(request):
    """Добавление модели рации."""
    if request.method == 'POST':
        form = RadioModelForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Модель рации успешно добавлена!')
            return redirect('radio_models_list')
    else:
        form = RadioModelForm()
    
    context = {
        'form': form,
        'page_title': 'Добавить модель рации',
    }
    return render(request, 'tasks/radio_model_form.html', context)


@login_required
def radio_model_edit(request, pk):
    """Редактирование модели рации."""
    radio_model = get_object_or_404(RadioModel, pk=pk)
    
    if request.method == 'POST':
        form = RadioModelForm(request.POST, instance=radio_model)
        if form.is_valid():
            form.save()
            messages.success(request, 'Модель рации успешно обновлена!')
            return redirect('radio_models_list')
    else:
        form = RadioModelForm(instance=radio_model)
    
    context = {
        'form': form,
        'page_title': 'Редактировать модель рации',
        'radio_model': radio_model,
    }
    return render(request, 'tasks/radio_model_form.html', context)


@login_required
def radio_list(request):
    """Список всех раций."""
    radios = RadioType.objects.all().select_related('model').order_by('model__name', 'sn')
    active_count = radios.filter(is_active=True).count()
    
    context = {
        'radios': radios,
        'active_count': active_count,
        'page_title': 'Справочник раций',
    }
    return render(request, 'tasks/radio_list.html', context)


@login_required
def radio_add(request):
    """Добавление рации в справочник."""
    if request.method == 'POST':
        form = RadioTypeForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Рация успешно добавлена в справочник!')
            return redirect('radio_list')
    else:
        form = RadioTypeForm()
    
    context = {
        'form': form,
        'page_title': 'Добавить рацию',
    }
    return render(request, 'tasks/radio_form.html', context)


@login_required
def radio_edit(request, pk):
    """Редактирование рации в справочнике."""
    radio = get_object_or_404(RadioType, pk=pk)
    
    if request.method == 'POST':
        form = RadioTypeForm(request.POST, instance=radio)
        if form.is_valid():
            form.save()
            messages.success(request, 'Рация успешно обновлена!')
            return redirect('radio_list')
    else:
        form = RadioTypeForm(instance=radio)
    
    context = {
        'form': form,
        'page_title': 'Редактировать рацию',
        'radio': radio,
    }
    return render(request, 'tasks/radio_form.html', context)


@login_required
def radio_assign(request):
    """Выдача рации."""
    if request.method == 'POST':
        form = RadioAssignmentForm(request.POST, request.FILES)
        if form.is_valid():
            assignment = form.save(commit=False)
            assignment.assignee_name = form.cleaned_data.get('assignee_name')
            assignment.issued_by = request.user
            assignment.save()
            messages.success(request, f'Рация {assignment.device} выдана {assignment.assignee_name}!')
            return redirect('radio_assignments_list')
    else:
        form = RadioAssignmentForm()
    
    context = {
        'form': form,
        'page_title': 'Выдача рации',
        'assignment': None,  # Явно указываем None для новой выдачи
    }
    return render(request, 'tasks/radio_assign_form.html', context)


@login_required
def radio_assignment_edit(request, pk):
    """Редактирование выдачи рации."""
    assignment = get_object_or_404(RadioAssignment, pk=pk)
    
    if request.method == 'POST':
        form = RadioAssignmentForm(request.POST, request.FILES, instance=assignment)
        if form.is_valid():
            assignment = form.save(commit=False)
            assignment.assignee_name = form.cleaned_data.get('assignee_name')
            # issued_by не меняем
            assignment.save()
            messages.success(request, 'Выдача рации успешно обновлена!')
            return redirect('radio_assignments_list')
    else:
        form = RadioAssignmentForm(instance=assignment)
        # Устанавливаем начальное значение для поля assignee_name
        form.fields['assignee_name'].initial = assignment.assignee_name
    
    context = {
        'form': form,
        'page_title': 'Редактировать выдачу рации',
        'assignment': assignment,
    }
    return render(request, 'tasks/radio_assign_form.html', context)


@login_required
def radio_assignments_list(request):
    """Список выданных раций."""
    assignments = RadioAssignment.objects.all().order_by('-issued_at')
    active_assignments = assignments.filter(returned_at__isnull=True)
    
    context = {
        'assignments': assignments,
        'active_assignments': active_assignments,
        'page_title': 'Журнал выдачи раций',
    }
    return render(request, 'tasks/radio_assignments_list.html', context)


@login_required
def radio_return(request, pk):
    """Возврат рации."""
    assignment = get_object_or_404(RadioAssignment, pk=pk)
    
    if assignment.returned_at:
        messages.warning(request, 'Эта рация уже возвращена!')
        return redirect('radio_assignments_list')
    
    if request.method == 'POST':
        assignment.returned_at = timezone.now()
        assignment.save()
        messages.success(request, f'Рация {assignment.device} возвращена!')
        return redirect('radio_assignments_list')
    
    context = {
        'assignment': assignment,
        'page_title': 'Возврат рации',
    }
    return render(request, 'tasks/radio_return.html', context)


@login_required
def radio_repair(request):
    """Учет ремонта рации."""
    if request.method == 'POST':
        form = RadioRepairForm(request.POST)
        if form.is_valid():
            repair = form.save(commit=False)
            repair.performed_by = request.user
            repair.save()
            messages.success(request, f'Ремонт рации {repair.device} зарегистрирован!')
            return redirect('radio_repairs_list')
    else:
        form = RadioRepairForm()
    
    context = {
        'form': form,
        'page_title': 'Учет ремонта рации',
    }
    return render(request, 'tasks/radio_repair_form.html', context)


@login_required
def radio_repairs_list(request):
    """Список ремонтов раций."""
    repairs = RadioRepair.objects.all().order_by('-repair_date')
    
    context = {
        'repairs': repairs,
        'page_title': 'История ремонтов раций',
    }
    return render(request, 'tasks/radio_repairs_list.html', context)


@login_required
def radio_assignment_act_excel(request, pk: int):
    """Экспорт акта приема-передачи оборудования в Excel по конкретной выдаче."""
    assignment = get_object_or_404(RadioAssignment.objects.select_related('device__model', 'issued_by'), pk=pk)
    device = assignment.device

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Акт передачи'

    # Определяем стили границ
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # Заголовок
    ws.merge_cells('A1:E1')
    ws['A1'] = 'Акт приёма-передачи оборудования'
    ws['A1'].font = Font(size=14, bold=True)
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30

    # Подзаголовок
    ws.merge_cells('A3:E3')
    ws['A3'] = 'Настоящий акт подтверждает передачу следующего оборудования:'

    # Шапка таблицы
    headers = ['№', 'Наименование', 'Модель', 'Серийный номер', 'Состояние']
    header_row = 5
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=title)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border

    # Данные строки
    state_text = 'Исправен' if device.is_active else 'Не исправен'
    data_row = 6
    ws.cell(row=data_row, column=1, value=1)
    ws.cell(row=data_row, column=2, value='Рация')
    ws.cell(row=data_row, column=3, value=device.model.name)
    ws.cell(row=data_row, column=4, value=device.sn)
    ws.cell(row=data_row, column=5, value=state_text)
    
    # Применяем границы к ячейкам с данными
    for col in range(1, 6):
        cell = ws.cell(row=data_row, column=col)
        cell.border = thin_border
        cell.alignment = Alignment(vertical='center')
    
    ws.row_dimensions[data_row].height = 40

    # Комплектация
    ws.merge_cells('A8:E8')
    ws['A8'] = 'Комплектация: '
    
    # Примечание с текстом по шаблону
    ws.merge_cells('A10:E13')
    note_text = (
        'Примечание:\n\n'
        '1.\tС момента подписания настоящего Акта ответственность за сохранность оборудования несёт Принимающая сторона.\n\n'
        '2.\tПередача оборудования осуществляется для служебного использования.'
    )
    ws['A10'] = note_text
    ws['A10'].alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[10].height = 60

    # Подписи - Передал
    signature_row = 15
    ws.merge_cells(start_row=signature_row, start_column=1, end_row=signature_row, end_column=3)
    ws.cell(row=signature_row, column=1, value='Передал')
    ws.cell(row=signature_row, column=1).font = Font(bold=True)
    ws.cell(row=signature_row, column=1).alignment = Alignment(horizontal='left', vertical='center')
    
    # Подписи - Принял
    ws.merge_cells(start_row=signature_row, start_column=4, end_row=signature_row, end_column=5)
    ws.cell(row=signature_row, column=4, value='Принял')
    ws.cell(row=signature_row, column=4).font = Font(bold=True)
    ws.cell(row=signature_row, column=4).alignment = Alignment(horizontal='left', vertical='center')
    
    # Поля для заполнения - Передал
    fields_row = signature_row + 1
    # Участок
    ws.merge_cells(start_row=fields_row, start_column=1, end_row=fields_row, end_column=3)
    ws.cell(row=fields_row, column=1, value='Участок:')
    ws.cell(row=fields_row, column=1).alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[fields_row].height = 20
    
    fields_row += 1
    ws.merge_cells(start_row=fields_row, start_column=1, end_row=fields_row, end_column=3)
    ws.cell(row=fields_row, column=1, value='')
    ws.row_dimensions[fields_row].height = 20
    
    # Должность /ФИО
    fields_row += 1
    ws.merge_cells(start_row=fields_row, start_column=1, end_row=fields_row, end_column=3)
    ws.cell(row=fields_row, column=1, value='Должность /ФИО:')
    ws.cell(row=fields_row, column=1).alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[fields_row].height = 20
    
    fields_row += 1
    ws.merge_cells(start_row=fields_row, start_column=1, end_row=fields_row, end_column=3)
    ws.cell(row=fields_row, column=1, value='')
    ws.row_dimensions[fields_row].height = 20
    
    # Пустая строка для подписи
    fields_row += 1
    ws.merge_cells(start_row=fields_row, start_column=1, end_row=fields_row, end_column=3)
    ws.cell(row=fields_row, column=1, value='')
    ws.row_dimensions[fields_row].height = 20
    
    # Дата/Подпись
    fields_row += 1
    ws.merge_cells(start_row=fields_row, start_column=1, end_row=fields_row, end_column=3)
    ws.cell(row=fields_row, column=1, value='Дата/Подпись:')
    ws.cell(row=fields_row, column=1).alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[fields_row].height = 20
    
    # Поля для заполнения - Принял
    fields_row = signature_row + 1
    # Участок
    ws.merge_cells(start_row=fields_row, start_column=4, end_row=fields_row, end_column=5)
    ws.cell(row=fields_row, column=4, value='Участок:')
    ws.cell(row=fields_row, column=4).alignment = Alignment(horizontal='left', vertical='center')
    
    fields_row += 1
    ws.merge_cells(start_row=fields_row, start_column=4, end_row=fields_row, end_column=5)
    ws.cell(row=fields_row, column=4, value='')
    
    # Должность /ФИО
    fields_row += 1
    ws.merge_cells(start_row=fields_row, start_column=4, end_row=fields_row, end_column=5)
    ws.cell(row=fields_row, column=4, value='Должность /ФИО:')
    ws.cell(row=fields_row, column=4).alignment = Alignment(horizontal='left', vertical='center')
    
    fields_row += 1
    ws.merge_cells(start_row=fields_row, start_column=4, end_row=fields_row, end_column=5)
    ws.cell(row=fields_row, column=4, value='')
    
    # Пустая строка для подписи
    fields_row += 1
    ws.merge_cells(start_row=fields_row, start_column=4, end_row=fields_row, end_column=5)
    ws.cell(row=fields_row, column=4, value='')
    
    # Дата/Подпись
    fields_row += 1
    ws.merge_cells(start_row=fields_row, start_column=4, end_row=fields_row, end_column=5)
    ws.cell(row=fields_row, column=4, value='Дата/Подпись:')
    ws.cell(row=fields_row, column=4).alignment = Alignment(horizontal='left', vertical='center')
    
    # Применяем границы ко всем ячейкам блока подписей
    for row in range(signature_row, signature_row + 7):
        for col in range(1, 6):
            cell = ws.cell(row=row, column=col)
            cell.border = thin_border

    # Ширины столбцов (оптимизированы для A4 portrait)
    # Передал: A+B+C = 10+15+15 = 40
    # Принял: D+E = 20+20 = 40
    widths = [10, 15, 15, 20, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Настройки печати для A4 (вертикальный формат)
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    # Ответ
    file_name = f"act_assignment_{assignment.id}.xlsx"
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{file_name}"'
    wb.save(response)
    return response
    # --- Добавить в конец views.py ---
# Не забудь добавить импорт вверху views.py:
# from .models import IdleRecord, ModelTagConfig
# from .clickhouse_client import ClickHouseClient
# from collections import defaultdict

@login_required
def telemetry_monitor(request):
    """Страница мониторинга телеметрии работающей техники."""
    from .models import IdleRecord, ModelTagConfig
    from .clickhouse_client import ClickHouseClient
    from collections import defaultdict
    
    hours = int(request.GET.get('hours', 1))
    show_ok = request.GET.get('show_ok', '') == '1'
    model_filter = request.GET.get('model', '')
    
    results = []
    stats = {'ok': 0, 'warn': 0, 'fail': 0, 'total_objects': 0, 'working': 0}
    ch_connected = False
    error_message = ''
    
    try:
        # 1. Проверяем ClickHouse
        ch = ClickHouseClient()
        ch_connected = ch.test_connection()
        
        if not ch_connected:
            error_message = 'Нет подключения к ClickHouse'
            raise Exception(error_message)
        
        # 2. Активные конфигурации тегов
        tag_configs = ModelTagConfig.objects.filter(
            is_active=True
        ).select_related('oes_model')
        
        if not tag_configs:
            error_message = 'Нет активных конфигураций тегов. Настройте их через Django Admin.'
            raise Exception(error_message)
        
        # Группируем по (модель, таблица)
        configs_by_model_table = defaultdict(list)
        models_with_configs = set()
        for tc in tag_configs:
            key = (tc.oes_model_id, tc.clickhouse_table)
            configs_by_model_table[key].append(tc)
            models_with_configs.add(tc.oes_model_id)
        
        # 3. Объекты с мониторингом
        all_objects = OesObject.objects.filter(
            model_id__in=models_with_configs,
            source_id__isnull=False
        ).select_related('model', 'model__category')
        
        # Фильтр по модели
        if model_filter:
            all_objects = all_objects.filter(model_id=model_filter)
        
        stats['total_objects'] = all_objects.count()
        
        # 4. ID объектов в простое
        idle_object_ids = set(
            IdleRecord.objects.filter(
                end_dt__isnull=True,
                oes_object__isnull=False
            ).values_list('oes_object__source_id', flat=True)
        )
        
        # 5. Работающие машины
        working_objects = [
            obj for obj in all_objects
            if obj.source_id not in idle_object_ids
        ]
        stats['working'] = len(working_objects)
        
        # 6. Группируем по (модель, таблица)
        objects_by_model_table = defaultdict(list)
        for obj in working_objects:
            for key in configs_by_model_table:
                model_id, ch_table = key
                if obj.model_id == model_id:
                    objects_by_model_table[key].append(obj)
        
        # 7. Запросы в ClickHouse
        for (model_id, ch_table), objects in objects_by_model_table.items():
            if not objects:
                continue
            
            tag_confs = configs_by_model_table[(model_id, ch_table)]
            object_uuids = [obj.mdm_object_uuid for obj in objects if obj.mdm_object_uuid]
            
            if not object_uuids:
                continue
            
            ch_results = ch.check_tags_batch(
                table=ch_table,
                object_uuids=object_uuids,
                tag_configs=tag_confs,
                hours=hours
            )
            
            obj_by_uuid = {obj.mdm_object_uuid: obj for obj in objects if obj.mdm_object_uuid}
            
            for uuid, tags_data in ch_results.items():
                obj = obj_by_uuid.get(uuid)
                if not obj:
                    continue
                
                has_fail = any(d['status'] == 'fail' for d in tags_data.values())
                has_warn = any(d['status'] == 'warn' for d in tags_data.values())
                
                if has_fail:
                    overall = 'fail'
                elif has_warn:
                    overall = 'warn'
                else:
                    overall = 'ok'
                
                # Считаем статистику
                for d in tags_data.values():
                    stats[d['status']] += 1
                
                # Добавляем в результаты (если show_ok или есть проблемы)
                if show_ok or overall != 'ok':
                    results.append({
                        'object': obj,
                        'overall': overall,
                        'tags': tags_data,
                    })
        
        # Сортировка: fail первые, потом warn, потом ok
        priority_order = {'fail': 0, 'warn': 1, 'ok': 2}
        results.sort(key=lambda r: (priority_order.get(r['overall'], 3), r['object'].name))
    
    except Exception as e:
        if not error_message:
            error_message = str(e)
    
    # Модели для фильтра
    monitored_models = OesModel.objects.filter(
        tag_configs__is_active=True
    ).distinct().order_by('name')
    
    context = {
        'results': results,
        'stats': stats,
        'ch_connected': ch_connected,
        'error_message': error_message,
        'hours': hours,
        'show_ok': show_ok,
        'model_filter': model_filter,
        'monitored_models': monitored_models,
        'page_title': 'Мониторинг телеметрии',
    }
    return render(request, 'tasks/telemetry_monitor.html', context)
# --- Добавить в конец views.py (после telemetry_monitor) ---

@login_required
def create_telemetry_task(request, object_id):
    """Создание задачи из мониторинга телеметрии с защитой от дублей."""
    from .models import ModelTagConfig
    
    oes_object = get_object_or_404(OesObject, pk=object_id)
    
    # Проверяем дедупликацию: есть ли открытая задача по этому объекту
    # с источником 'TELEMETRY'
    existing_task = Task.objects.filter(
        oes_object=oes_object,
        external_source='TELEMETRY',
    ).exclude(
        status=Task.STATUS_DONE
    ).first()
    
    if existing_task:
        messages.warning(
            request, 
            f'Задача #{existing_task.id} по объекту {oes_object.name} уже существует!'
        )
        return redirect('task_detail', pk=existing_task.pk)
    
    # Создаём задачу
    model_name = oes_object.model.name if oes_object.model else 'Неизвестная модель'
    
    task = Task.objects.create(
        title=f'Нет телеметрии: {oes_object.name} ({model_name})',
        description=(
            f'Автоматически создана из мониторинга телеметрии.\n'
            f'Объект: {oes_object.name}\n'
            f'Модель: {model_name}\n'
            f'Проблема: отсутствие или некорректные данные телеметрии.\n'
            f'Дата обнаружения: {timezone.now().strftime("%d.%m.%Y %H:%M")}'
        ),
        status=Task.STATUS_WAITING,
        priority=Task.PRIORITY_MEDIUM,
        oes_object=oes_object,
        reporter=request.user,
        external_source='TELEMETRY',
        external_id=f'telem_{oes_object.id}_{timezone.now().strftime("%Y%m%d")}',
    )
    
    messages.success(request, f'Задача #{task.id} создана для {oes_object.name}!')
    return redirect('telemetry_monitor')


@login_required
def create_telemetry_tasks_bulk(request):
    """Массовое создание задач для всех проблемных объектов."""
    if request.method != 'POST':
        return redirect('telemetry_monitor')
    
    from .models import IdleRecord, ModelTagConfig
    from .clickhouse_client import ClickHouseClient
    from collections import defaultdict
    
    hours = int(request.POST.get('hours', 1))
    created_count = 0
    skipped_count = 0
    
    try:
        ch = ClickHouseClient()
        if not ch.test_connection():
            messages.error(request, 'Нет подключения к ClickHouse')
            return redirect('telemetry_monitor')
        
        # Повторяем логику из telemetry_monitor
        tag_configs = ModelTagConfig.objects.filter(is_active=True).select_related('oes_model')
        
        configs_by_model_table = defaultdict(list)
        models_with_configs = set()
        for tc in tag_configs:
            key = (tc.oes_model_id, tc.clickhouse_table)
            configs_by_model_table[key].append(tc)
            models_with_configs.add(tc.oes_model_id)
        
        all_objects = OesObject.objects.filter(
            model_id__in=models_with_configs,
            source_id__isnull=False
        ).select_related('model')
        
        idle_object_ids = set(
            IdleRecord.objects.filter(
                end_dt__isnull=True,
                oes_object__isnull=False
            ).values_list('oes_object__source_id', flat=True)
        )
        
        working_objects = [
            obj for obj in all_objects
            if obj.source_id not in idle_object_ids
        ]
        
        # Объекты с уже открытыми задачами TELEMETRY
        existing_task_object_ids = set(
            Task.objects.filter(
                external_source='TELEMETRY'
            ).exclude(
                status=Task.STATUS_DONE
            ).values_list('oes_object_id', flat=True)
        )
        
        objects_by_model_table = defaultdict(list)
        for obj in working_objects:
            for key in configs_by_model_table:
                model_id, ch_table = key
                if obj.model_id == model_id:
                    objects_by_model_table[key].append(obj)
        
        for (model_id, ch_table), objects in objects_by_model_table.items():
            if not objects:
                continue
            
            tag_confs = configs_by_model_table[(model_id, ch_table)]
            object_uuids = [obj.mdm_object_uuid for obj in objects if obj.mdm_object_uuid]
            
            if not object_uuids:
                continue
            
            ch_results = ch.check_tags_batch(
                table=ch_table,
                object_uuids=object_uuids,
                tag_configs=tag_confs,
                hours=hours
            )
            
            obj_by_uuid = {obj.mdm_object_uuid: obj for obj in objects if obj.mdm_object_uuid}
            
            for uuid, tags_data in ch_results.items():
                obj = obj_by_uuid.get(uuid)
                if not obj:
                    continue
                
                has_problems = any(
                    d['status'] in ('fail', 'warn')
                    for d in tags_data.values()
                )
                
                if not has_problems:
                    continue
                
                # Дедупликация
                if obj.id in existing_task_object_ids:
                    skipped_count += 1
                    continue
                
                # Собираем описание проблем
                problem_details = []
                max_priority = Task.PRIORITY_LOW
                priority_order = {
                    'critical': 0, 'high': 1, 'medium': 2, 'low': 3
                }
                
                for tag_name, tag_data in tags_data.items():
                    if tag_data['status'] in ('fail', 'warn'):
                        problem_details.append(f"  • {tag_name}: {tag_data['detail']}")
                        # Берём максимальный приоритет из конфигов
                        for tc in tag_confs:
                            if tc.tag_name == tag_name:
                                if priority_order.get(tc.task_priority, 3) < priority_order.get(max_priority, 3):
                                    max_priority = tc.task_priority
                
                model_name = obj.model.name if obj.model else 'Неизвестная модель'
                
                Task.objects.create(
                    title=f'Нет телеметрии: {obj.name} ({model_name})',
                    description=(
                        f'Автоматически создана из мониторинга телеметрии.\n'
                        f'Объект: {obj.name}\n'
                        f'Модель: {model_name}\n\n'
                        f'Обнаруженные проблемы:\n'
                        f'{chr(10).join(problem_details)}\n\n'
                        f'Дата обнаружения: {timezone.now().strftime("%d.%m.%Y %H:%M")}'
                    ),
                    status=Task.STATUS_WAITING,
                    priority=max_priority,
                    oes_object=obj,
                    reporter=request.user,
                    external_source='TELEMETRY',
                    external_id=f'telem_{obj.id}_{timezone.now().strftime("%Y%m%d")}',
                )
                created_count += 1
                existing_task_object_ids.add(obj.id)
        
        messages.success(
            request,
            f'Создано задач: {created_count}, пропущено (уже есть): {skipped_count}'
        )
    
    except Exception as e:
        messages.error(request, f'Ошибка: {str(e)}')
    
    return redirect('telemetry_monitor')